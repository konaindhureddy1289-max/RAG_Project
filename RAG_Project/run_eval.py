"""
Runs the 20 test questions through your RAG assistant and records
retrieved sources, the answer and latency into the spreadsheet.

Setup (inside your activated venv):
    pip install requests openpyxl

Usage:
    1. Start your API:  uvicorn main:app --reload
    2. Put rag_eval_test_set.xlsx in the same folder as this script
    3. Edit the CONFIG section below to match YOUR endpoint
    4. Run:  python run_eval.py
    5. Open rag_eval_results.xlsx and fill Hit? / Accurate? / Grounded?
"""
import time
import requests
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ---------------- CONFIG: edit these to match your API ----------------
URL = "http://127.0.0.1:8000/ask"   # your endpoint
QUESTION_FIELD = "query"             # the JSON field your API expects
INPUT_FILE = "rag_eval_test_set.xlsx"
OUTPUT_FILE = "rag_eval_results.xlsx"
TIMEOUT = 60


def extract_answer(data):
    """Pull the answer text out of your API response. Edit if needed."""
    if isinstance(data, str):
        return data
    for key in ("answer", "response", "result", "output"):
        if key in data:
            return str(data[key])
    return str(data)


def extract_sources(data):
    """Pull retrieved chunks/sources out of the response. Edit if needed."""
    if not isinstance(data, dict):
        return ""
    for key in ("sources", "retrieved_chunks", "chunks", "context", "documents"):
        if key in data:
            val = data[key]
            if isinstance(val, list):
                return " | ".join(str(v)[:150] for v in val)
            return str(val)[:450]
    return "(no sources field found: edit extract_sources)"
# ----------------------------------------------------------------------


def main():
    wb = load_workbook(INPUT_FILE)
    ws = wb["Test Set"]

    # Add an answer column (M) next to the existing ones
    ws["M1"] = "Assistant answer"
    ws["M1"].font = Font(name="Arial", bold=True, color="FFFFFF")
    ws["M1"].fill = PatternFill("solid", fgColor="1F3864")
    ws.column_dimensions["M"].width = 60

    for row in range(2, ws.max_row + 1):
        test_id = ws.cell(row, 1).value
        question = ws.cell(row, 3).value
        if not question:
            continue

        start = time.perf_counter()
        try:
            resp = requests.post(URL, json={QUESTION_FIELD: question}, timeout=TIMEOUT)
            latency = time.perf_counter() - start
            resp.raise_for_status()
            try:
                data = resp.json()
            except ValueError:
                data = resp.text
            answer = extract_answer(data)
            sources = extract_sources(data)
        except Exception as e:  # keep going if one question fails
            latency = time.perf_counter() - start
            answer, sources = f"ERROR: {e}", ""

        ws.cell(row, 7).value = sources            # G: retrieved sources
        ws.cell(row, 11).value = round(latency, 2)  # K: latency (s)
        c = ws.cell(row, 13, answer)                # M: answer
        c.font = Font(name="Arial", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        print(f"{test_id}: {latency:.2f}s  {answer[:80]!r}")

    wb.save(OUTPUT_FILE)
    print(f"\nDone. Open {OUTPUT_FILE} and score Hit? / Accurate? / Grounded?")


if __name__ == "__main__":
    main()
